"""
extractor_datos.py  –  MÓDULO 1 (I/O con la planilla Excel). NO importa Streamlit.

Responsabilidades (ver ARQUITECTURA_Y_REGLAS.md §3 y §4.1):
  * Leer de 'Hoja 2' la lista de cursos (col. A) y de distritos (col. C), sin hardcodear nada.
  * Priorizar los distritos de la Región 10 (067, 057, 063, 070, 044, 093, 101).
  * Guardar una fila evaluada en 'MERITUADOS' o 'NO MERITUADOS' desde la fila 9
    (primera fila con la columna C vacía), sin alterar estilos, celdas combinadas,
    validaciones de datos, comentarios ni fórmulas del libro.
  * Ofrecer los bytes del Excel actualizado para `st.download_button`.

El archivo original NUNCA se modifica: se trabaja sobre `salida/PLANILLA_TRABAJO.xlsx`.
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

# --------------------------------------------------------------------------- #
# Constantes
# --------------------------------------------------------------------------- #
BASE_DIR = Path(__file__).resolve().parent
RUTA_ORIGINAL = BASE_DIR / "PLANILLA DE DATOS DE ASPIRANTES.xlsx"
DIR_SALIDA = BASE_DIR / "salida"
RUTA_TRABAJO = DIR_SALIDA / "PLANILLA_TRABAJO.xlsx"
DIR_BACKUP = DIR_SALIDA / "backup"
MAX_BACKUPS = 20

FILA_INICIAL = 9
HOJA_MER = "MERITUADOS"
HOJA_NO_MER = "NO MERITUADOS"
HOJA_LISTAS = "Hoja 2"
COL_CURSOS = 1      # Hoja 2, columna A
COL_DISTRITOS = 3   # Hoja 2, columna C
COL_DNI = 3         # columna C en ambas hojas de datos
COL_DISTRITO_EVALUADOR = "D1"

# Diccionario de Familias Profesionales (Catálogo 2027, §13.1)
FAMILIAS_PROFESIONALES: dict[str, str] = {
    "00": "00 - Cargos Docentes (Pañol / Preceptor)",
    "AA": "AA - Actividades Agropecuarias",
    "AC": "AC - Administración y Comercialización",
    "AE": "AE - Aeronáutica",
    "AP": "AP - Actividades Pesqueras",
    "CC": "CC - Cuero y Calzado",
    "CS": "CS - Comunicación Social",
    "EE": "EE - Energía Eléctrica",
    "EM": "EM - Electromecánica",
    "ET": "ET - Electrónica",
    "GH": "GH - Gastronomía / Hotelería y Turismo",
    "IA": "IA - Industria Alimentaria",
    "IC": "IC - Industrias Culturales y Creativas",
    "IG": "IG - Industria Gráfica",
    "IM": "IM - Informática",
    "LT": "LT - Logística y Transporte",
    "MA": "MA - Automotriz",
    "MC": "MC - Módulos Comunes",
    "MM": "MM - Madera y Mueble",
    "MT": "MT - Metalmecánica",
    "OA": "OA - Oficios Artesanales",
    "OC": "OC - Construcciones",
    "PA": "PA - Industria Papelera",
    "PL": "PL - Industria Plástica",
    "SC": "SC - Cuidados y Servicios Sociocomunitarios",
    "SM": "SM - Seguridad y Medio Ambiente",
    "SP": "SP - Estética Profesional",
    "SS": "SS - Salud",
    "TC": "TC - Telecomunicaciones",
    "TI": "TI - Textil e Indumentaria",
    "OT": "OT - Otros / Nuevos Cursos",
}

RUTA_CURSOS_PERSONALIZADOS = DIR_SALIDA / "cursos_personalizados.json"

CURSOS_CATALOGO_2027_INICIALES = [
    "AA54 - Armadora/or de Invernáculos",
    "AA55 - Podadora/or Frutícola",
    "AC36 - Promotora/or de cooperativas y mutuales",
    "IA19 - Elaboradora/or de productos lácteos diversificados",
    "OC49 - Operaria/o Ladrillera/o",
    "IG10 - Operadora/or en Diseño para la Industria Gráfica",
    "MM12 - Desarrolladora/or de Mobiliario y Productos de Carpintería",
    "MT44 - Montadora/or Cañista Industrial",
    "SC03 - Operadora/or de Cuidados de Personas con Discapacidad",
    "TC01 - Empalmadora/or de Fibra Óptica",
    "TC02 - Instaladora/or Reparadora/or de redes de acceso FTTH y redes internas del cliente",
    "TC03 - Analista en Redes de Telecomunicaciones",
    "TC04 - Operadora/or de Redes de Fibra Óptica",
    "MC01 - Orientación Profesional",
]

# Códigos de distrito (Región 10) a mostrar primero, en este orden:
# 067-MARCOS PAZ, 057-GENERAL LAS HERAS, 063-LUJÁN, 070-MERCEDES,
# 044-GENERAL RODRÍGUEZ, 093-SAN ANDRÉS DE GILES, 101-SUIPACHA.
# Solo son CÓDIGOS: los nombres completos se leen de la planilla.
CODIGOS_PRIORITARIOS_REG10 = ("067", "057", "063", "070", "044", "093", "101")

# Mapas clave -> letra de columna (ARQUITECTURA §3.1 y §3.2)
COL_MER = {
    "orden": "A", "listado": "B", "dni": "C", "apellido": "D", "nombres": "E",
    "especialidad": "F",
    "dist1": "G", "pts_dist1": "H", "dist2": "I", "pts_dist2": "J",
    "dist3": "K", "pts_dist3": "L",
    "fi": "M",
    "fe_sec_orient": "N", "fe_sec_tecnico": "O", "fe_superior": "P", "fe_universitario": "Q",
    "cap_docente": "R", "cap_especifica": "S",
    "desemp_anios": "T", "desemp_pts": "U",
    "form_sup_fp": "V", "antig_docencia": "W", "antig_fp": "X",
    "calif1": "Y", "calif2": "Z", "art59": "AA", "puntaje": "AB",
}
COL_NO_MER = {
    "orden": "A", "dni": "C", "apellido": "D", "nombres": "E", "especialidad": "F",
    "dist1": "G", "dist2": "H", "dist3": "I", "motivo": "J", "observaciones": "K",
}  # 'listado' (B) queda vacío por diseño (decisión D4)

REQUERIDOS_MER = ("dni", "apellido", "nombres", "especialidad", "dist1", "listado", "fi", "puntaje")
REQUERIDOS_NO_MER = ("dni", "apellido", "nombres", "especialidad", "dist1", "motivo")

_RE_REG10 = re.compile(r"^\s*Reg\s+10\s+Distrito:\s*(\d+)\s*-", re.IGNORECASE)

# Verificación de estructura: (celda, texto con el que debe empezar, en mayúsculas)
_CHEQUEOS_MER = (("A4", "ORDEN"), ("B4", "LISTADO"), ("C5", "DNI"), ("D5", "APELLIDO"),
                 ("E5", "NOMBRES"), ("M4", "INSTRUCTOR"), ("R4", "CAPACITA"),
                 ("AA4", "ART 59"), ("AB4", "PUNTAJE"))
_CHEQUEOS_NO_MER = (("A4", "ORDEN"), ("C5", "DNI"), ("D5", "APELLIDO"), ("E5", "NOMBRES"),
                    ("J4", "NO MERITUA"), ("K4", "OBSERVACIONES"))


# --------------------------------------------------------------------------- #
# Excepciones
# --------------------------------------------------------------------------- #
class ErrorEstructuraPlanilla(Exception):
    """La planilla no tiene la estructura esperada (hojas, encabezados o listas)."""


class PlanillaBloqueada(Exception):
    """No se puede escribir el archivo (normalmente está abierto en Excel)."""


# --------------------------------------------------------------------------- #
# Utilidades internas
# --------------------------------------------------------------------------- #
def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _col_idx(letra: str) -> int:
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letra)


def _lista_limpia(valores: Iterable) -> list[str]:
    """Sin vacíos, `strip()` y sin duplicados, conservando el orden original."""
    vistos: dict[str, None] = {}
    for v in valores:
        t = _texto(v)
        if t:
            vistos.setdefault(t, None)
    return list(vistos)


def _normalizar_dni(dni) -> int:
    solo_digitos = re.sub(r"\D", "", str(dni))
    if not solo_digitos:
        raise ValueError(f"DNI inválido: {dni!r}")
    return int(solo_digitos)


def _normalizar_valor(valor):
    if isinstance(valor, float):
        return round(valor, 2)
    if isinstance(valor, str):
        return valor.strip()
    return valor


def _validar_estructura(wb) -> None:
    for hoja in (HOJA_MER, HOJA_NO_MER, HOJA_LISTAS):
        if hoja not in wb.sheetnames:
            raise ErrorEstructuraPlanilla(f"Falta la hoja '{hoja}' en la planilla.")
    for hoja, chequeos in ((HOJA_MER, _CHEQUEOS_MER), (HOJA_NO_MER, _CHEQUEOS_NO_MER)):
        ws = wb[hoja]
        for celda, inicio in chequeos:
            if not _texto(ws[celda].value).upper().startswith(inicio):
                raise ErrorEstructuraPlanilla(
                    f"Hoja '{hoja}', celda {celda}: se esperaba un encabezado que empiece con "
                    f"'{inicio}' y se encontró {ws[celda].value!r}. La planilla cambió de estructura."
                )


def _validar_registro(registro: dict, requeridos: tuple[str, ...]) -> None:
    faltan = [k for k in requeridos if registro.get(k) in (None, "")]
    if faltan:
        raise ValueError(f"Faltan datos obligatorios en el registro: {', '.join(faltan)}")


def _hacer_backup(ruta: Path) -> None:
    try:
        DIR_BACKUP.mkdir(parents=True, exist_ok=True)
        marca = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(ruta, DIR_BACKUP / f"{ruta.stem}_{marca}{ruta.suffix}")
        backups = sorted(DIR_BACKUP.glob(f"{ruta.stem}_*{ruta.suffix}"))
        for viejo in backups[:-MAX_BACKUPS]:
            viejo.unlink(missing_ok=True)
    except PermissionError as exc:
        raise PlanillaBloqueada(f"No se pudo hacer el backup de {ruta.name}: {exc}") from exc


def _guardar_atomico(wb, ruta: Path) -> None:
    """Guarda en un temporal y reemplaza; si el destino está bloqueado, avisa sin dañarlo."""
    tmp = ruta.with_name(ruta.stem + ".tmp" + ruta.suffix)
    try:
        wb.save(tmp)
        os.replace(tmp, ruta)
    except PermissionError as exc:
        tmp.unlink(missing_ok=True)
        raise PlanillaBloqueada(
            f"No se pudo guardar '{ruta.name}'. Cerrá el archivo en Excel y reintentá."
        ) from exc


# --------------------------------------------------------------------------- #
# Archivo de trabajo
# --------------------------------------------------------------------------- #
def asegurar_copia_trabajo() -> Path:
    """Crea `salida/PLANILLA_TRABAJO.xlsx` copiando el original si todavía no existe."""
    if not RUTA_TRABAJO.exists():
        if not RUTA_ORIGINAL.exists():
            raise FileNotFoundError(f"No se encontró la planilla original: {RUTA_ORIGINAL}")
        DIR_SALIDA.mkdir(parents=True, exist_ok=True)
        shutil.copy2(RUTA_ORIGINAL, RUTA_TRABAJO)
    return RUTA_TRABAJO


# --------------------------------------------------------------------------- #
# Lectura de listas y datos generales
# --------------------------------------------------------------------------- #
def es_distrito_region10(distrito: str) -> bool:
    return bool(_RE_REG10.match(distrito or ""))


def ordenar_distritos(distritos: list[str]) -> list[str]:
    """Región 10 primero (los prioritarios en el orden de CODIGOS_PRIORITARIOS_REG10,
    luego el resto de la Región 10) y después las demás regiones, en su orden original."""
    por_codigo: dict[str, str] = {}
    resto_reg10: list[str] = []
    otras: list[str] = []
    for d in distritos:
        m = _RE_REG10.match(d)
        if not m:
            otras.append(d)
        elif m.group(1) in CODIGOS_PRIORITARIOS_REG10 and m.group(1) not in por_codigo:
            por_codigo[m.group(1)] = d
        else:
            resto_reg10.append(d)
    prioritarios = [por_codigo[c] for c in CODIGOS_PRIORITARIOS_REG10 if c in por_codigo]
    return prioritarios + resto_reg10 + otras


# --------------------------------------------------------------------------- #
# Cursos Personalizados y Familias Profesionales (§13)
# --------------------------------------------------------------------------- #
def inicializar_cursos_personalizados() -> None:
    """Crea salida/cursos_personalizados.json con los 14 cursos del Catálogo 2027 si no existe."""
    if not RUTA_CURSOS_PERSONALIZADOS.exists():
        DIR_SALIDA.mkdir(parents=True, exist_ok=True)
        payload = {"version": 1, "cursos": list(CURSOS_CATALOGO_2027_INICIALES)}
        RUTA_CURSOS_PERSONALIZADOS.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8"
        )


def obtener_cursos_personalizados() -> list[str]:
    """Devuelve la lista limpia de cursos guardados en salida/cursos_personalizados.json."""
    inicializar_cursos_personalizados()
    try:
        data = json.loads(RUTA_CURSOS_PERSONALIZADOS.read_text(encoding="utf-8"))
        raw = data.get("cursos", [])
        return _lista_limpia(raw)
    except Exception:
        return list(CURSOS_CATALOGO_2027_INICIALES)


def agregar_o_editar_curso_personalizado(
    codigo_nombre_nuevo: str,
    codigo_nombre_anterior: str | None = None
) -> tuple[bool, str]:
    """Agrega un curso nuevo o reemplaza uno anterior si se editó.
    Guarda en salida/cursos_personalizados.json y retorna (True, mensaje) o (False, error)."""
    nuevo = (codigo_nombre_nuevo or "").strip()
    if not nuevo or len(nuevo) < 3:
        return False, "El nombre o código del curso debe tener al menos 3 caracteres."

    cursos = obtener_cursos_personalizados()
    if codigo_nombre_anterior and codigo_nombre_anterior.strip() in cursos:
        idx = cursos.index(codigo_nombre_anterior.strip())
        cursos[idx] = nuevo
        msg = f"Curso '{codigo_nombre_anterior}' modificado a '{nuevo}' con éxito."
    else:
        if nuevo in cursos:
            return True, f"El curso '{nuevo}' ya existía en la base de datos."
        cursos.append(nuevo)
        msg = f"Curso '{nuevo}' guardado en la base de datos con éxito."

    try:
        DIR_SALIDA.mkdir(parents=True, exist_ok=True)
        payload = {"version": 1, "cursos": cursos}
        RUTA_CURSOS_PERSONALIZADOS.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8"
        )
        return True, msg
    except Exception as exc:
        return False, f"Error al guardar el curso: {exc}"


def obtener_prefijo_familia(curso: str | None) -> str:
    """Devuelve el prefijo de 2 letras del curso (ej: 'IM30 - ...' -> 'IM').
    Si no coincide con ninguna familia conocida, devuelve 'OT'."""
    if not curso or len(str(curso).strip()) < 2:
        return "OT"
    pref = str(curso).strip()[:2].upper()
    return pref if pref in FAMILIAS_PROFESIONALES else "OT"


def obtener_nombre_familia(curso: str | None) -> str:
    """Devuelve el nombre completo descriptivo de la familia."""
    pref = obtener_prefijo_familia(curso)
    return FAMILIAS_PROFESIONALES.get(pref, FAMILIAS_PROFESIONALES["OT"])


def obtener_cursos_por_familia(prefijo: str, cursos: list[str]) -> list[str]:
    """Filtra una lista de cursos por prefijo de familia o 'TODAS LAS FAMILIAS'."""
    if not prefijo or prefijo in ("TODAS", "TODAS LAS FAMILIAS"):
        return list(cursos)
    p = prefijo.strip().upper()
    if " - " in p:
        p = p.split(" - ")[0].strip()
    if p not in FAMILIAS_PROFESIONALES:
        p = "OT"
    return [c for c in cursos if obtener_prefijo_familia(c) == p]


def obtener_todos_los_cursos(ruta_base: Path = RUTA_ORIGINAL) -> list[str]:
    """Une los cursos de 'Hoja 2' (columna A) con los de cursos_personalizados.json,
    sin duplicados y ordenados alfabéticamente."""
    cursos_base: list[str] = []
    p_base = Path(ruta_base)
    if p_base.exists():
        try:
            wb = load_workbook(p_base, read_only=True, data_only=True)
            if HOJA_LISTAS in wb.sheetnames:
                ws = wb[HOJA_LISTAS]
                cursos_base = _lista_limpia(
                    v for (v,) in ws.iter_rows(min_col=COL_CURSOS, max_col=COL_CURSOS, values_only=True)
                )
            wb.close()
        except Exception:
            pass

    personalizados = obtener_cursos_personalizados()
    todos = _lista_limpia(cursos_base + personalizados)
    todos.sort()
    return todos


def cargar_listas(ruta: Path = RUTA_ORIGINAL) -> tuple[list[str], list[str]]:
    """(cursos, distritos) leídos de 'Hoja 2' (col. A y col. C) unificados con cursos personalizados.
    Sin vacíos ni duplicados. Los distritos salen ordenados con la Región 10 primero."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if HOJA_LISTAS not in wb.sheetnames:
            raise ErrorEstructuraPlanilla(f"Falta la hoja '{HOJA_LISTAS}'.")
        ws = wb[HOJA_LISTAS]
        distritos = _lista_limpia(v for (v,) in ws.iter_rows(
            min_col=COL_DISTRITOS, max_col=COL_DISTRITOS, values_only=True))
    finally:
        wb.close()
    if not distritos:
        raise ErrorEstructuraPlanilla("'Hoja 2', columna C: no se encontraron distritos.")
    cursos = obtener_todos_los_cursos(ruta)
    if not cursos:
        raise ErrorEstructuraPlanilla("'Hoja 2', columna A: no se encontraron cursos.")
    return cursos, ordenar_distritos(distritos)


def leer_distrito_evaluador(ruta: Path = RUTA_ORIGINAL) -> str:
    """Valor de MERITUADOS!D1 (ej. 'Reg 10 Distrito: 067-MARCOS PAZ')."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb[HOJA_MER]
        fila = next(ws.iter_rows(min_row=1, max_row=1, min_col=4, max_col=4, values_only=True))
        return _texto(fila[0])
    finally:
        wb.close()


# --------------------------------------------------------------------------- #
# Búsqueda de fila libre y consultas
# --------------------------------------------------------------------------- #
def siguiente_fila_libre(ws, letras_destino: Iterable[str] | None = None) -> int:
    """Primera fila >= 9 con la columna C vacía (y sin celdas combinadas en las columnas destino)."""
    idx_destino = [_col_idx(c) for c in (letras_destino or [])]
    fila = FILA_INICIAL
    while True:
        if _texto(ws.cell(fila, COL_DNI).value) == "":
            if not any(isinstance(ws.cell(fila, i), MergedCell) for i in idx_destino):
                return fila
        fila += 1


def siguiente_orden(ws) -> int:
    """max(col. A numérica desde la fila 9) + 1, o 1 si no hay datos."""
    mayor = 0
    for fila in range(FILA_INICIAL, ws.max_row + 1):
        v = ws.cell(fila, 1).value
        try:
            mayor = max(mayor, int(float(v)))
        except (TypeError, ValueError):
            continue
    return mayor + 1


def buscar_por_dni(ruta: Path, dni) -> list[dict]:
    """Registros ya cargados con ese DNI en ambas hojas:
    [{'hoja', 'fila', 'especialidad', 'puntaje'|'motivo'}]. Sirve para avisar duplicados."""
    objetivo = _normalizar_dni(dni)
    encontrados: list[dict] = []
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        for hoja, ultima_col, extra in ((HOJA_MER, 28, "puntaje"), (HOJA_NO_MER, 11, "motivo")):
            ws = wb[hoja]
            filas = ws.iter_rows(min_row=FILA_INICIAL, max_col=ultima_col, values_only=True)
            for n, fila in enumerate(filas, start=FILA_INICIAL):
                fila = tuple(fila or ()) + (None,) * (28 - len(fila or ()))
                if fila[2] in (None, ""):
                    continue
                try:
                    if _normalizar_dni(fila[2]) != objetivo:
                        continue
                except ValueError:
                    continue
                valor_extra = fila[27] if extra == "puntaje" else fila[9]
                encontrados.append({"hoja": hoja, "fila": n,
                                    "especialidad": _texto(fila[5]), extra: valor_extra})
    finally:
        wb.close()
    return encontrados


def resumen_cargados(ruta: Path = RUTA_TRABAJO) -> dict:
    """{'meritudos': n, 'no_meritudos': n} (filas con DNI cargado desde la fila 9)."""
    if not Path(ruta).exists():
        return {"meritudos": 0, "no_meritudos": 0}
    cuenta = {}
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        for clave, hoja in (("meritudos", HOJA_MER), ("no_meritudos", HOJA_NO_MER)):
            ws = wb[hoja]
            cuenta[clave] = sum(
                1 for (v,) in ws.iter_rows(min_row=FILA_INICIAL, min_col=COL_DNI,
                                           max_col=COL_DNI, values_only=True)
                if _texto(v) != "")
    finally:
        wb.close()
    return cuenta


# --------------------------------------------------------------------------- #
# Escritura
# --------------------------------------------------------------------------- #
def _preparar_registro(registro: dict, requeridos: tuple[str, ...]) -> dict:
    """Valida campos obligatorios y normaliza DNI (int), apellido y nombres (MAYÚSCULAS)."""
    registro = dict(registro)
    _validar_registro(registro, requeridos)
    registro["dni"] = _normalizar_dni(registro["dni"])
    for k in ("apellido", "nombres"):
        registro[k] = " ".join(_texto(registro[k]).upper().split())
    return registro


def _escribir_fila(ws, registro: dict, mapa: dict[str, str]) -> int:
    """Escribe `registro` en la primera fila libre de `ws` (solo .value). Devuelve la fila."""
    fila = siguiente_fila_libre(ws, mapa.values())
    orden = siguiente_orden(ws)
    for clave, letra in mapa.items():
        if clave == "orden":
            valor = orden
        elif clave in registro:
            valor = _normalizar_valor(registro[clave])
        else:
            continue
        ws[f"{letra}{fila}"].value = valor   # solo .value: no se tocan estilos
    return fila


def _guardar_fila(registro: dict, ruta: Path, hoja: str, mapa: dict[str, str],
                  requeridos: tuple[str, ...]) -> int:
    registro = _preparar_registro(registro, requeridos)

    ruta = Path(ruta)
    if ruta == RUTA_TRABAJO:
        asegurar_copia_trabajo()
    _hacer_backup(ruta)

    # Carga completa (sin read_only) para conservar estilos, merges, validaciones y fórmulas.
    wb = load_workbook(ruta)
    try:
        _validar_estructura(wb)
        fila = _escribir_fila(wb[hoja], registro, mapa)
        _guardar_atomico(wb, ruta)
    finally:
        wb.close()
    return fila


def guardar_meritudo(registro: dict, ruta: Path = RUTA_TRABAJO) -> int:
    """Escribe el registro en MERITUADOS (A..AB, desde la fila 9). Devuelve el nº de fila."""
    return _guardar_fila(registro, ruta, HOJA_MER, COL_MER, REQUERIDOS_MER)


def guardar_no_meritudo(registro: dict, ruta: Path = RUTA_TRABAJO) -> int:
    """Escribe el registro en NO MERITUADOS (A, C..K, desde la fila 9). Devuelve el nº de fila."""
    return _guardar_fila(registro, ruta, HOJA_NO_MER, COL_NO_MER, REQUERIDOS_NO_MER)


def exportar_excel_bytes(meritudos: list[dict], no_meritudos: list[dict],
                         distrito_evaluador: str | None = None,
                         ruta_base: Path = RUTA_ORIGINAL) -> bytes:
    """Inserta TODOS los registros dados sobre una copia en memoria de `ruta_base` (que no se
    modifica) y devuelve el .xlsx resultante. Respeta formato, celdas combinadas, validaciones
    y fórmulas. Si se indica `distrito_evaluador`, se escribe en D1 de ambas hojas.
    Además, agrega cualquier curso personalizado o usado a 'Hoja 2' (Col A) y actualiza
    los rangos de validación de datos para que Excel reconozca todos los cursos (§13.2.3)."""
    wb = load_workbook(ruta_base)
    try:
        _validar_estructura(wb)
        if distrito_evaluador:
            for hoja in (HOJA_MER, HOJA_NO_MER):
                celda = wb[hoja][COL_DISTRITO_EVALUADOR]
                if not isinstance(celda, MergedCell):
                    celda.value = distrito_evaluador

        # §13.2.3: Inserción de cursos personalizados o usados en 'Hoja 2' Col A
        if HOJA_LISTAS in wb.sheetnames:
            ws2 = wb[HOJA_LISTAS]
            cursos_existentes = set()
            ultima_fila = 0
            for r in range(1, ws2.max_row + 1):
                val = ws2.cell(r, COL_CURSOS).value
                if val is not None and str(val).strip():
                    cursos_existentes.add(str(val).strip())
                    ultima_fila = r

            # Cursos a incorporar (de registros cargados y de cursos_personalizados.json)
            cursos_nuevos = []
            for reg in meritudos + no_meritudos:
                esp = reg.get("especialidad") or reg.get("curso")
                if esp and str(esp).strip() and str(esp).strip() not in cursos_existentes and str(esp).strip() not in cursos_nuevos:
                    cursos_nuevos.append(str(esp).strip())
            for c in obtener_cursos_personalizados():
                c_limpio = str(c).strip()
                if c_limpio and c_limpio not in cursos_existentes and c_limpio not in cursos_nuevos:
                    cursos_nuevos.append(c_limpio)

            if cursos_nuevos:
                for c in cursos_nuevos:
                    ultima_fila += 1
                    ws2.cell(ultima_fila, COL_CURSOS, value=c)
                    cursos_existentes.add(c)

            nuevo_rango_dv = f"'{HOJA_LISTAS}'!$A$1:$A${ultima_fila}"
            for hoja_nombre in (HOJA_MER, HOJA_NO_MER):
                ws_h = wb[hoja_nombre]
                for dv in ws_h.data_validations.dataValidation:
                    if dv.formula1 and "Hoja 2" in str(dv.formula1) and "$A$" in str(dv.formula1):
                        dv.formula1 = nuevo_rango_dv

        for reg in meritudos:
            _escribir_fila(wb[HOJA_MER], _preparar_registro(reg, REQUERIDOS_MER), COL_MER)
        for reg in no_meritudos:
            _escribir_fila(wb[HOJA_NO_MER], _preparar_registro(reg, REQUERIDOS_NO_MER), COL_NO_MER)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    finally:
        wb.close()


# --------------------------------------------------------------------------- #
# Descarga
# --------------------------------------------------------------------------- #
def obtener_bytes_excel(ruta: Path = RUTA_TRABAJO) -> bytes:
    """Contenido del Excel actualizado, para `st.download_button(data=...)`."""
    if Path(ruta) == RUTA_TRABAJO:
        asegurar_copia_trabajo()
    return Path(ruta).read_bytes()


def nombre_archivo_descarga(ruta: Path = RUTA_TRABAJO) -> str:
    """Nombre sugerido: PLANILLA_<distrito evaluador>_<AAAAMMDD_HHMM>.xlsx"""
    try:
        distrito = leer_distrito_evaluador(ruta)
    except Exception:  # noqa: BLE001 - el nombre es solo cosmético
        distrito = ""
    seguro = re.sub(r"[^A-Za-z0-9]+", "_", distrito).strip("_") or "REG10"
    return f"PLANILLA_{seguro}_{datetime.now():%Y%m%d_%H%M}.xlsx"


if __name__ == "__main__":  # verificación rápida desde consola
    cursos_, distritos_ = cargar_listas()
    print(f"Distrito evaluador: {leer_distrito_evaluador()}")
    print(f"{len(cursos_)} cursos, {len(distritos_)} distritos")
    print("Primeros distritos:", *distritos_[:8], sep="\n  ")
